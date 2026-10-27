"""
app/api/reportes.py
-------------------
Resumen para la pantalla de inicio y reportes de gasto ejercido.
"""

from datetime import date
from io import BytesIO
from typing import Optional

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse

from app.auth import usuario_actual, requiere_rol
from app.database import supabase
from app.utils import credito

router = APIRouter(prefix="/reportes", tags=["Reportes"])

# Estados en los que la orden ya representa gasto ejercido.
ESTADOS_GASTO = ("pagada", "recoleccion", "comprobacion", "cerrada")

ROLES_FINANZAS = ("admin", "compras", "pagos")

SELECT_GASTO = "*,proveedor:proveedor_id(id,nombre,rfc)"

# Una sola línea: PostgREST pierde los embebidos si la cadena es multilínea.
# El proveedor va con * para que dias_credito llegue sin listarlo aquí.
SELECT_CREDITO = (
    "id,folio,fecha,estado,total,monto_comprobado,proveedor_id,proveedor:proveedor_id(*),"
    "pagos(fecha_pago,monto,metodo,referencia),recolecciones(completado,fecha_real),"
    "facturas(serie,folio,fecha_emision)"
)

ETIQUETAS_PAGO = {
    "transferencia":   "Transferencia",
    "tarjeta_credito": "Tarjeta de crédito",
    "tarjeta_debito":  "Tarjeta de débito",
    "efectivo":        "Efectivo",
    "credito":         "Crédito del proveedor",
}

ETIQUETAS_ESTADO = {
    "pagada":      "Pagada",
    "recoleccion":  "Por recolectar",
    "comprobacion": "Por comprobar",
    "cerrada":      "Cerrada",
}


def _rango_mes(hoy: date) -> tuple[str, str]:
    primero = hoy.replace(day=1)
    if hoy.month == 12:
        siguiente = date(hoy.year + 1, 1, 1)
    else:
        siguiente = date(hoy.year, hoy.month + 1, 1)
    return primero.isoformat(), siguiente.isoformat()


def _contar(estado: str) -> int:
    res = (
        supabase.table("ordenes_compra")
        .select("id", count="exact")
        .eq("estado", estado)
        .execute()
    )
    return res.count or 0


@router.get("/dashboard")
def dashboard(usuario=Depends(usuario_actual)):
    """Cifras de la pantalla de inicio.

    Los montos solo se devuelven a los roles que ya ven Control de gastos.
    """
    inicio_mes, inicio_siguiente = _rango_mes(date.today())

    pagos_mes = (
        supabase.table("pagos")
        .select("monto, orden_id")
        .gte("fecha_pago", inicio_mes)
        .lt("fecha_pago", inicio_siguiente)
        .execute()
    ).data or []

    resumen = {
        "en_autorizacion":  _contar("autorizacion"),
        "por_recolectar":   _contar("recoleccion"),
        "por_comprobar":    _contar("comprobacion"),
        "pagadas_mes":      len({p["orden_id"] for p in pagos_mes}),
        "mes":              inicio_mes,
        "incluye_montos":   usuario["rol"] in ROLES_FINANZAS,
    }
    if resumen["incluye_montos"]:
        resumen["gasto_mes"] = round(sum(float(p["monto"] or 0) for p in pagos_mes), 2)
        resumen["facturas_pendientes"] = (
            supabase.table("facturas").select("id", count="exact")
            .is_("orden_id", "null").execute()
        ).count or 0
        cuentas = _estados_de_cuenta()["totales"]
        resumen["credito_vencido"] = cuentas["vencido"]
        resumen["credito_ordenes_vencidas"] = cuentas["ordenes_vencidas"]
    return resumen


def _estados_de_cuenta() -> dict:
    ordenes = (
        supabase.table("ordenes_compra")
        .select(SELECT_CREDITO)
        .eq("tipo_pago", "credito")
        .in_("estado", list(credito.ESTADOS_DEUDA))
        .execute()
    ).data or []
    return credito.estados_de_cuenta(ordenes, date.today())


@router.get("/credito")
def estados_de_cuenta_credito(usuario=Depends(requiere_rol(*ROLES_FINANZAS))):
    """Saldo y vencimientos por proveedor de las compras a crédito."""
    return _estados_de_cuenta()


def _consultar_gastos(
    desde: Optional[str],
    hasta: Optional[str],
    proveedor_id: Optional[str],
    tipo_pago: Optional[str],
) -> list[dict]:
    query = (
        supabase.table("ordenes_compra")
        .select(SELECT_GASTO)
        .in_("estado", list(ESTADOS_GASTO))
        .order("fecha", desc=True)
    )
    if desde:
        query = query.gte("fecha", desde)
    if hasta:
        query = query.lte("fecha", hasta)
    if proveedor_id:
        query = query.eq("proveedor_id", proveedor_id)
    if tipo_pago:
        if tipo_pago not in ETIQUETAS_PAGO:
            raise HTTPException(status_code=400, detail=f"Tipo de pago inválido: {tipo_pago}")
        query = query.eq("tipo_pago", tipo_pago)
    return query.execute().data or []


@router.get("/gastos")
def listar_gastos(
    desde: Optional[str] = Query(None, description="Fecha de orden mínima (YYYY-MM-DD)"),
    hasta: Optional[str] = Query(None, description="Fecha de orden máxima (YYYY-MM-DD)"),
    proveedor_id: Optional[str] = Query(None),
    tipo_pago: Optional[str] = Query(None),
    usuario=Depends(requiere_rol(*ROLES_FINANZAS)),
):
    return _consultar_gastos(desde, hasta, proveedor_id, tipo_pago)


@router.get("/gastos/excel")
def exportar_gastos_excel(
    desde: Optional[str] = Query(None),
    hasta: Optional[str] = Query(None),
    proveedor_id: Optional[str] = Query(None),
    tipo_pago: Optional[str] = Query(None),
    usuario=Depends(requiere_rol(*ROLES_FINANZAS)),
):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    ordenes = _consultar_gastos(desde, hasta, proveedor_id, tipo_pago)

    wb = Workbook()
    ws = wb.active
    ws.title = "Gasto ejercido"

    encabezados = ["Folio", "Fecha", "Proveedor", "RFC", "Tipo de pago",
                   "Estado", "Subtotal", "IVA", "Total"]
    ws.append(encabezados)

    relleno = PatternFill("solid", fgColor="F4F1EA")
    for celda in ws[1]:
        celda.font = Font(bold=True, color="26241D")
        celda.fill = relleno
        celda.alignment = Alignment(horizontal="center")

    for o in ordenes:
        prov = o.get("proveedor") or {}
        ws.append([
            o.get("folio"),
            o.get("fecha"),
            prov.get("nombre") or "—",
            prov.get("rfc") or "—",
            ETIQUETAS_PAGO.get(o.get("tipo_pago"), o.get("tipo_pago")),
            ETIQUETAS_ESTADO.get(o.get("estado"), o.get("estado")),
            float(o.get("subtotal") or 0),
            float(o.get("iva") or 0),
            float(o.get("total") or 0),
        ])

    ultima = ws.max_row
    if ordenes:
        ws.cell(row=ultima + 1, column=6, value="Total").font = Font(bold=True)
        for col in (7, 8, 9):
            letra = get_column_letter(col)
            celda = ws.cell(row=ultima + 1, column=col,
                            value=f"=SUM({letra}2:{letra}{ultima})")
            celda.font = Font(bold=True)

    for col in range(7, 10):
        for fila in range(2, ws.max_row + 1):
            ws.cell(row=fila, column=col).number_format = '"$"#,##0.00'

    anchos = [8, 12, 42, 16, 18, 16, 14, 12, 14]
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = "A2"

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    nombre = f"gasto-ejercido-{date.today().isoformat()}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{nombre}"'},
    )
